import io
import pytest
from docx import Document
from httpx import AsyncClient, ASGITransport
from app.main import app


@pytest.mark.asyncio
async def test_parse_jd_docx_endpoint():
    # Create sample docx document in memory
    doc = Document()
    doc.add_heading("Lead Systems Architect", level=1)
    doc.add_paragraph("We are hiring a Lead Systems Architect with experience in Rust, C++, and distributed storage.")
    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        files = {
            "file": ("architect_jd.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        }
        res = await client.post("/api/v1/roles/parse-jd", files=files)
        assert res.status_code == 200
        body = res.json()
        assert "data" in body
        data = body["data"]
        assert "Lead Systems Architect" in data["text"]
        assert "distributed storage" in data["text"]
        assert data["suggested_title"] == "Lead Systems Architect"
        # Ensure no mock text or attachment wrapper appears in text
        assert "Senior software engineer with expertise in distributed systems" not in data["text"]
        assert "[Attached Document:" not in data["text"]


@pytest.mark.asyncio
async def test_parse_jd_txt_endpoint():
    txt_content = b"# Frontend Team Lead\nLooking for React, Next.js, and TypeScript expert."
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        files = {"file": ("frontend_lead.txt", txt_content, "text/plain")}
        res = await client.post("/api/v1/roles/parse-jd", files=files)
        assert res.status_code == 200
        data = res.json()["data"]
        assert "Frontend Team Lead" in data["text"]
        assert "TypeScript" in data["text"]
        assert data["suggested_title"] == "Frontend Team Lead"


@pytest.mark.asyncio
async def test_parse_jd_empty_rejection():
    empty_content = b"   "
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        files = {"file": ("empty.txt", empty_content, "text/plain")}
        res = await client.post("/api/v1/roles/parse-jd", files=files)
        assert res.status_code == 400
